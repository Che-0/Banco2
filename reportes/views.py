from django.shortcuts import render

# Create your views here.
from django.shortcuts import render
from django.contrib.auth.decorators import login_required, user_passes_test
from django.http import HttpResponse
from django.db.models import Count, Sum, Q
from django.utils import timezone
from datetime import timedelta
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

from clientes.models import Cliente, CuentaBancaria, Transferencia
from accounts.models import User

def es_admin_o_empleado(user):
    return user.is_authenticated and (user.es_admin or user.es_empleado or user.is_superuser)

@login_required
@user_passes_test(es_admin_o_empleado)
def reportes_dashboard(request):
    """Página principal de reportes"""
    hoy = timezone.now().date()
    hace_30_dias = hoy - timedelta(days=30)

    context = {
        'total_clientes': Cliente.objects.count(),
        'clientes_activos': Cliente.objects.filter(estado='ACTIVO').count(),
        'total_cuentas': CuentaBancaria.objects.count(),
        'saldo_total': CuentaBancaria.objects.aggregate(total=Sum('saldo'))['total'] or 0,
        'total_transferencias': Transferencia.objects.count(),
        'transferencias_mes': Transferencia.objects.filter(fecha__date__gte=hace_30_dias).count(),
        'monto_transferido_mes': Transferencia.objects.filter(
            fecha__date__gte=hace_30_dias
        ).aggregate(total=Sum('monto'))['total'] or 0,
        'usuarios_por_rol': User.objects.values('rol').annotate(total=Count('id')),
    }
    return render(request, 'reportes/dashboard_reportes.html', context)

@login_required
@user_passes_test(es_admin_o_empleado)
def reporte_clientes(request):
    """Listado de clientes con filtros"""
    clientes = Cliente.objects.select_related('usuario', 'creado_por').all()

    estado = request.GET.get('estado')
    if estado:
        clientes = clientes.filter(estado=estado)

    q = request.GET.get('q')
    if q:
        clientes = clientes.filter(
            Q(nombres__icontains=q) |
            Q(apellidos__icontains=q) |
            Q(numero_documento__icontains=q)
        )

    context = {
        'clientes': clientes,
        'estado_seleccionado': estado,
        'q': q or '',
    }
    return render(request, 'reportes/reporte_clientes.html', context)

@login_required
@user_passes_test(es_admin_o_empleado)
def exportar_clientes_excel(request):
    """Exportar clientes a Excel"""
    clientes = Cliente.objects.all().order_by('-fecha_registro')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clientes"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0D6EFD")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    headers = [
        "Nombres", "Apellidos", "Documento", "Teléfono",
        "Email", "Ciudad", "Estado", "Fecha Registro"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row, cliente in enumerate(clientes, 2):
        ws.cell(row=row, column=1, value=cliente.nombres).border = thin_border
        ws.cell(row=row, column=2, value=cliente.apellidos).border = thin_border
        ws.cell(row=row, column=3, value=cliente.numero_documento).border = thin_border
        ws.cell(row=row, column=4, value=cliente.telefono or "").border = thin_border
        ws.cell(row=row, column=5, value=cliente.email or "").border = thin_border
        ws.cell(row=row, column=6, value=cliente.ciudad or "").border = thin_border
        ws.cell(row=row, column=7, value=cliente.get_estado_display()).border = thin_border
        ws.cell(row=row, column=8, value=cliente.fecha_registro.strftime("%d/%m/%Y")).border = thin_border

    # Ajustar anchos
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column].width = max_length + 2

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=clientes.xlsx'
    wb.save(response)
    return response

@login_required
@user_passes_test(es_admin_o_empleado)
def reporte_transferencias(request):
    """Reporte de transferencias"""
    transferencias = Transferencia.objects.select_related(
        'cuenta_origen__cliente',
        'cuenta_destino__cliente',
        'realizado_por'
    ).all()[:100]  # últimas 100

    context = {
        'transferencias': transferencias,
    }
    return render(request, 'reportes/reporte_transferencias.html', context)
